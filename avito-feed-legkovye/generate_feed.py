"""Собирает fleet-legkovye.xlsx для Авито Автозагрузки (категория Автомобили -> Новые)
из всех active JSON-записей в data/. Схема полей вытащена реальным запросом к
autoload/v1/user-docs/node/{slug}/fields (узел 'novye'), не гаданием."""
import json
import glob
import openpyxl

REQUIRED_IN_PRACTICE = ['Цена в валюте']

HEADERS = ['Уникальный идентификатор объявления', 'Начало размещения', 'Окончание размещения',
           'Способ размещения', 'Услуга продвижения', 'Номер объявления на Авито', 'Контактное лицо',
           'Номер телефона', 'Адрес', 'Широта', 'Долгота', 'Способ связи', 'Интернет звонки',
           'Устройства для приёма звонков', 'Тип автомобиля', 'Статус наличия', 'Категория',
           'Описание объявления', 'Цена', 'Цвет', 'Названия фото', 'Ссылки на фото', 'Ссылка на видео',
           'Марка', 'Модель', 'Поколение', 'Модификация', 'Комплектация', 'Тип двигателя',
           'Коробка передач', 'Объём двигателя', 'Год выпуска', 'Количество дверей', 'Тип кузова',
           'Привод', 'Мощность', 'Руль', 'Салон', 'Цвет салона', 'VIN или номер кузова', 'Цена в валюте',
           'Валюта', 'Управление климатом', 'Управление климатом (дополнительные опции)',
           'Салон (дополнительные опции)', 'Обогрев', 'Электропривод', 'Память настроек',
           'Помощь при вождении', 'Мультимедиа и навигация', 'Фары']

# Avito ожидает несколько значений checkbox-поля через запятую в одной ячейке —
# формат подтверждён только по общей документации Автозагрузки, не проверен
# живой синхронизацией (см. reference/avito_legkovye_raw_catalog/ и заметку в памяти).
CHECKBOX_SEP = ", "

RAW_BASE = "https://raw.githubusercontent.com/aj123123-hub/belkar-avito-feed/main/avito-feed-legkovye/"


def build_row(rec):
    photo_urls = " | ".join(RAW_BASE + p for p in rec.get("photos", []))
    row = {
        'Уникальный идентификатор объявления': rec["id"],
        'Контактное лицо': rec.get("manager_name", ""),
        'Номер телефона': rec.get("phone", ""),
        'Адрес': rec["address"],
        'Способ связи': 'По телефону и в сообщениях',
        'Тип автомобиля': 'Новые',
        'Статус наличия': rec["availability"],
        'Категория': 'Автомобили',
        'Описание объявления': rec["description"],
        'Цена': rec["price"],
        'Цена в валюте': rec["price"],
        'Валюта': rec["currency"],
        'Цвет': rec["color"],
        'Ссылки на фото': photo_urls,
        'Марка': rec["make"],
        'Модель': rec["model"],
        'Поколение': rec["generation"],
        'Модификация': rec["modification"],
        'Комплектация': rec["complectation"],
        'Тип двигателя': rec["fuel_type"],
        'Коробка передач': rec["transmission"],
        'Объём двигателя': rec["engine_size_l"],
        'Год выпуска': rec["year"],
        'Количество дверей': rec["doors"],
        'Тип кузова': rec["body_type"],
        'Привод': rec["drive_type"],
        'Мощность': rec["power_hp"],
        'Салон': rec["interior"],
        'Цвет салона': rec["interior_color"],
        'VIN или номер кузова': rec["vin"],
        'Руль': rec.get("wheel_type", "Левый"),
        'Услуга продвижения': rec.get("promotion", ""),
        'Управление климатом': rec.get("climate_control", ""),
        'Управление климатом (дополнительные опции)': CHECKBOX_SEP.join(rec.get("climate_control_options", [])),
        'Салон (дополнительные опции)': CHECKBOX_SEP.join(rec.get("interior_options", [])),
        'Обогрев': CHECKBOX_SEP.join(rec.get("heating", [])),
        'Электропривод': CHECKBOX_SEP.join(rec.get("electric_drive", [])),
        'Память настроек': CHECKBOX_SEP.join(rec.get("memory_settings", [])),
        'Помощь при вождении': CHECKBOX_SEP.join(rec.get("driving_assistance", [])),
        'Мультимедиа и навигация': CHECKBOX_SEP.join(rec.get("multimedia", [])),
        'Фары': rec.get("lights", ""),
    }
    return [row.get(h, "") for h in HEADERS]


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Объявления"
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.number_format = "@"

    count = 0
    r = 2
    missing_required = []
    for path in sorted(glob.glob("data/*.json")):
        with open(path, encoding="utf-8") as f:
            rec = json.load(f)
        if rec.get("status") != "active":
            continue
        if not rec.get("vin"):
            missing_required.append(f"{path}: VIN не заполнен — Авито отклонит синхронизацию без VIN или номера кузова")
        values = build_row(rec)
        row_map = dict(zip(HEADERS, values))
        for field in REQUIRED_IN_PRACTICE:
            if field in row_map and not str(row_map[field]).strip():
                missing_required.append(f"{path}: пусто поле {field!r} (обязательно на практике)")
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.number_format = "@"
        r += 1
        count += 1

    if missing_required:
        raise SystemExit("Есть карточки с проблемами, фид не собран:\n" + "\n".join(missing_required))

    wb.save("fleet-legkovye.xlsx")
    print(f"Собрано {count} активных объявлений -> fleet-legkovye.xlsx")


if __name__ == "__main__":
    main()
