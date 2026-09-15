"""Собирает fleet-construction.xlsx для Авито Автозагрузки (категории «Экскаваторы»
гусеничные/колёсные + «Бульдозеры») из всех active JSON-записей в data/.

Три листа в одном файле — Авито поддерживает несколько категорий в одном источнике
через отдельные листы (подтверждено на связке Грузовики/Прицепы, см. память проекта).
"""
import json
import glob
import openpyxl

RAW_BASE = "https://raw.githubusercontent.com/aj123123-hub/belkar-avito-feed/main/avito-feed-construction/"

HEADERS_EXCAVATOR_TRACKED = [
    'Уникальный идентификатор объявления', 'Начало размещения', 'Окончание размещения',
    'Способ размещения', 'Услуга продвижения', 'Номер объявления на Авито', 'Контактное лицо',
    'Номер телефона', 'Адрес', 'Широта', 'Долгота', 'Способ связи', 'Описание объявления',
    'Категория', 'Цена', 'Зоны показа', 'Названия фото', 'Ссылки на фото', 'Ссылка на видео',
    'Название объявления', 'Интернет звонки', 'Устройства для приёма звонков', 'Вид техники',
    'Валюта', 'НДС включён', 'Утильсбор включён', 'Цена в валюте', 'Доступность',
    'Адрес стоянки', 'Скидка за лизинг', 'Скидка при покупке от двух единиц', 'Скидка от дилера',
    'Доставка', 'Установка дополнительного оборудования', 'Официальная гарантия',
    'Дополнительные условия гарантии', 'Подарки', 'URL видеофайла', 'Состояние',
    'ПТС или ПСМ', 'Моточасы', 'Тип техники', 'Марка', 'Модель', 'Мощность двигателя',
    'Эксплуатационная масса', 'Ширина гусеничной ленты', 'VIN, номер кузова или SN',
    'Год выпуска', 'Глубина копания', 'Грузоподъёмность погрузочного ковша',
    'Объём погрузочного ковша', 'Максимальная высота выгрузки', 'Объём ковша',
    'Дополнительная гидролиния', 'Тип управления копательным ковшом', 'TTL (Auction)', 'Цена (Auction)',
]

HEADERS_EXCAVATOR_WHEELED = [h for h in HEADERS_EXCAVATOR_TRACKED if h != 'Ширина гусеничной ленты']

HEADERS_BULLDOZER = [
    'Уникальный идентификатор объявления', 'Начало размещения', 'Окончание размещения',
    'Способ размещения', 'Услуга продвижения', 'Номер объявления на Авито', 'Контактное лицо',
    'Номер телефона', 'Адрес', 'Широта', 'Долгота', 'Способ связи', 'Описание объявления',
    'Категория', 'Цена', 'Зоны показа', 'Названия фото', 'Ссылки на фото', 'Ссылка на видео',
    'Интернет звонки', 'Устройства для приёма звонков', 'Название объявления', 'Вид техники',
    'Валюта', 'НДС включён', 'Утильсбор включён', 'Цена в валюте', 'Доступность',
    'Адрес стоянки', 'Скидка за лизинг', 'Скидка при покупке от двух единиц', 'Скидка от дилера',
    'Доставка', 'Установка дополнительного оборудования', 'Официальная гарантия',
    'Дополнительные условия гарантии', 'Подарки', 'URL видеофайла', 'Состояние',
    'ПТС или ПСМ', 'Моточасы', 'Марка', 'Модель', 'Тип техники', 'VIN, номер кузова или SN',
    'Год выпуска', 'Мощность кВт', 'Эксплуатационная масса бульдозера',
    'Максимальное заглубление', 'Максимальное тяговое усилие', 'Ширина отвала',
    'TTL (Auction)', 'Цена (Auction)',
]

SHEETS = {
    "gusenichnye_drugie": ("Гусеничные экскаваторы", HEADERS_EXCAVATOR_TRACKED),
    "kolyosnye_drugie": ("Колёсные экскаваторы", HEADERS_EXCAVATOR_WHEELED),
    "buldozery": ("Бульдозеры", HEADERS_BULLDOZER),
}


def photo_urls(rec):
    return " | ".join(RAW_BASE + p for p in rec.get("photos", []))


def build_row_excavator(rec, headers):
    row = {
        'Уникальный идентификатор объявления': rec["id"],
        'Контактное лицо': rec.get("manager_name", ""),
        'Номер телефона': rec.get("phone", ""),
        'Адрес': rec["address"],
        'Способ связи': 'По телефону и в сообщениях',
        'Описание объявления': rec["description"],
        'Категория': 'Грузовики и спецтехника',
        'Цена': rec["price"],
        'Цена в валюте': rec["price"],
        'Названия фото': "",
        'Ссылки на фото': photo_urls(rec),
        'Название объявления': rec.get("title", ""),
        'Вид техники': 'Экскаваторы',
        'Валюта': rec["currency"],
        'НДС включён': 'Да' if rec.get("vat_included") else 'Нет',
        'Доступность': rec["availability"],
        'Адрес стоянки': rec["address"],
        'Состояние': rec["condition"],
        'ПТС или ПСМ': rec["technical_passport"],
        'Тип техники': rec["type_of_vehicle"],
        'Марка': rec["make"],
        'Модель': rec["model"],
        'Мощность двигателя': rec["power_kw"],
        'Эксплуатационная масса': rec["operating_weight"],
        'Ширина гусеничной ленты': rec.get("track_width", ""),
        'VIN, номер кузова или SN': rec.get("vin", ""),
        'Год выпуска': rec["year"],
        'Глубина копания': rec.get("digging_depth", ""),
        'Объём ковша': rec.get("bucket_volume", ""),
    }
    return [row.get(h, "") for h in headers]


def build_row_bulldozer(rec, headers):
    row = {
        'Уникальный идентификатор объявления': rec["id"],
        'Контактное лицо': rec.get("manager_name", ""),
        'Номер телефона': rec.get("phone", ""),
        'Адрес': rec["address"],
        'Способ связи': 'По телефону и в сообщениях',
        'Описание объявления': rec["description"],
        'Категория': 'Грузовики и спецтехника',
        'Цена': rec["price"],
        'Цена в валюте': rec["price"],
        'Названия фото': "",
        'Ссылки на фото': photo_urls(rec),
        'Название объявления': rec.get("title", ""),
        'Вид техники': 'Бульдозеры',
        'Валюта': rec["currency"],
        'НДС включён': 'Да' if rec.get("vat_included") else 'Нет',
        'Доступность': rec["availability"],
        'Адрес стоянки': rec["address"],
        'Состояние': rec["condition"],
        'ПТС или ПСМ': rec["technical_passport"],
        'Марка': rec["make"],
        'Модель': rec["model"],
        'Тип техники': rec["type_of_vehicle"],
        'VIN, номер кузова или SN': rec.get("vin", ""),
        'Год выпуска': rec["year"],
        'Мощность кВт': rec["power_kw"],
        'Эксплуатационная масса бульдозера': rec.get("bulldozer_weight", ""),
        'Максимальное заглубление': rec.get("bulldozer_max_depth", ""),
        'Максимальное тяговое усилие': rec.get("bulldozer_pulling_force", ""),
        'Ширина отвала': rec.get("cutting_edge_width", ""),
    }
    return [row.get(h, "") for h in headers]


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    records_by_slug = {slug: [] for slug in SHEETS}
    for path in sorted(glob.glob("data/*.json")):
        with open(path, encoding="utf-8") as f:
            rec = json.load(f)
        if rec.get("status") != "active":
            continue
        vin = rec.get("vin", "")
        if vin and vin in rec.get("description", ""):
            raise SystemExit(f"{path}: VIN {vin} попал в текст описания — уберите")
        slug = rec.get("category_slug")
        if slug not in SHEETS:
            raise SystemExit(f"{path}: неизвестный category_slug {slug!r}")
        records_by_slug[slug].append((path, rec))

    total = 0
    for slug, (sheet_title, headers) in SHEETS.items():
        ws = wb.create_sheet(title=sheet_title)
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.number_format = "@"
        r = 2
        for path, rec in records_by_slug[slug]:
            if slug == "buldozery":
                values = build_row_bulldozer(rec, headers)
            else:
                values = build_row_excavator(rec, headers)
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=str(v) if v != "" else "")
                cell.number_format = "@"
            r += 1
            total += 1
        for col in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 40)
        print(f"{sheet_title}: {len(records_by_slug[slug])} объявлений")

    wb.save("fleet-construction.xlsx")
    print(f"\nfleet-construction.xlsx собран: {total} активных объявлений всего")


if __name__ == "__main__":
    main()
