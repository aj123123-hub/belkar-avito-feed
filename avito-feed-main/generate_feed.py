"""Собирает fleet-main.xlsx (аккаунт ООО БелКар, id 24645721) из active-записей data/*.json.

Формат — тот же, что у фида авитолога (английские имена полей Авито, один лист на все категории),
чтобы Авито сопоставил строки с уже существующими объявлениями по Id/AvitoId, а не создал новые.
"""
import concurrent.futures
import glob
import json
import subprocess
import sys

import openpyxl

SNAPSHOT = "reference/avitolog_feed_2026-09-14_snapshot.xlsx"
RAW_BASE = "https://raw.githubusercontent.com/aj123123-hub/belkar-avito-feed/main/avito-feed-main/"


def main():
    order = [h for h in next(openpyxl.load_workbook(SNAPSHOT).active.iter_rows(values_only=True)) if h]
    recs = []
    for path in sorted(glob.glob("data/*.json")):
        with open(path, encoding="utf-8") as f:
            rec = json.load(f)
        if rec.get("status") == "active":
            recs.append(rec)
    if len({r["id"] for r in recs}) != len(recs):
        raise SystemExit("повторяющийся Id в data/")

    used = {k for r in recs for k in r["fields"]} | {"ImageUrls"}
    headers = [h for h in order if h in used] + sorted(used - set(order))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Объявления"
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).number_format = "@"

    all_urls = set()
    for r_i, rec in enumerate(recs, 2):
        row = dict(rec["fields"])
        if "ПОДАРОК" in row.get("Description", ""):
            raise SystemExit(f"{rec['id']}: в описании остался блок с подарком")
        own = [RAW_BASE + p for p in rec["photos"]]
        all_urls.update(own)
        row["ImageUrls"] = " | ".join(own or rec.get("external_photo_urls", []))
        for c, h in enumerate(headers, 1):
            v = row.get(h, "")
            # требование Авито: все ячейки текстовым форматом
            ws.cell(row=r_i, column=c, value=str(v) if v != "" else "").number_format = "@"

    if "--skip-photo-check" not in sys.argv:
        print(f"Проверяю {len(all_urls)} ссылок на фото (curl)...")

        def check(u):
            return u, subprocess.run(["curl", "-s", "--max-time", "15", "-o", "/dev/null", "-w", "%{http_code}", u],
                                     capture_output=True, text=True).stdout

        with concurrent.futures.ThreadPoolExecutor(max_workers=20) as pool:
            bad = [f"{code} {u}" for u, code in pool.map(check, sorted(all_urls)) if code != "200"]
        if bad:
            raise SystemExit("Битые ссылки на фото (сначала запушьте фото):\n" + "\n".join(bad))
        print("Все фото отдают 200.")

    wb.save("fleet-main.xlsx")
    print(f"fleet-main.xlsx: {len(recs)} объявлений, {len(headers)} колонок")


if __name__ == "__main__":
    main()
